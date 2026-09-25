
# Schreibt die Messwerte (siehe measure.py) in Excel-Dateien:
#   Einzelner Versuch: eine Tabelle + Graph
#   Ganzer Ordner:     alles auf einem Blatt untereinander, je Versuch
#                      Tabelle + Graph, ganz unten der Durchschnitt mit
#                      Mittelwert +- Standardabweichung

from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import AreaChart, LineChart, Reference
from openpyxl.chart.legend import LegendEntry
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Font, PatternFill

HEADER_FILL = PatternFill("solid", fgColor="DDE7F0")
LINE_COLOR = "1F5A96"   # dunkelblau
BAND_COLOR = "BCD4EC"   # hellblau fuer die Standardabweichung

RUN_HEADER = [
    "Zeitpunkt", "Bild",
    "Abstand Mittelwert (px)", "Abstand Max (px)", "Zeile Max",
    "Fläche (px)", "Fläche relativ zu Zeitpunkt 1 (%)",
]
REL_COLUMN = 7  # Spalte G: Flaeche relativ
BAND_COLUMN = 9  # Durchschnitt: Spalten I/J sind die Hilfsspalten fuers Band

# Alles steht auf EINEM Blatt untereinander. Ein Graph ist ~9 cm hoch,
# das sind knapp 18 Excel-Zeilen - so viel Platz braucht jeder Block mindestens.
MIN_BLOCK_ROWS = 20


def relative_areas(measurements: list) -> list:
    start = measurements[0]["area"] if measurements else 0
    return [round(100.0 * m["area"] / start, 2) if start else None for m in measurements]


def write_header(ws, row: int, header: list) -> None:
    for col, text in enumerate(header, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 32


def set_widths(ws, widths: list) -> None:
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width


def style_axes(chart, n_points: int) -> None:
    chart.x_axis.title = "Zeitpunkt"
    chart.y_axis.title = "Fläche in % von Zeitpunkt 1"
    chart.y_axis.scaling.min = 0
    chart.y_axis.majorUnit = 20
    # openpyxl blendet die Achsen sonst in neueren Excel-Versionen aus
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.legend.position = "b"
    chart.height = 9
    chart.width = max(14, 3 * n_points)


def line_chart(ws, title: str, first_row: int, last_row: int) -> LineChart:
    #Ein Graph: relative Flaeche ueber die Zeitpunkte, startet bei 100 %
    chart = LineChart()
    chart.title = title
    data = Reference(ws, min_col=REL_COLUMN, min_row=first_row - 1, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first_row, max_row=last_row))

    series = chart.series[0]
    series.graphicalProperties.line.solidFill = LINE_COLOR
    series.graphicalProperties.line.width = 28575  # 2.25 pt
    series.marker = Marker(symbol="circle", size=7)
    series.marker.graphicalProperties.solidFill = LINE_COLOR
    series.marker.graphicalProperties.line.solidFill = LINE_COLOR
    series.smooth = False

    style_axes(chart, last_row - first_row + 1)
    chart.legend = None
    return chart


def write_run(ws, start: int, run_name: str, image_paths: list, measurements: list, rows: np.ndarray) -> int:
    #Ein Versuch ab Zeile start: Titel, Tabelle, Hinweise, Graph rechts daneben.
    #Gibt die erste freie Zeile nach dem Block zurueck.
    ws.cell(row=start, column=1, value=run_name).font = Font(bold=True, size=14)
    write_header(ws, start + 1, RUN_HEADER)

    first_row = start + 2
    for i, (path, m, rel) in enumerate(zip(image_paths, measurements, relative_areas(measurements))):
        values = [i + 1, Path(path).name, round(m["mean"], 2), m["max"], m["row_max"], m["area"], rel]
        for col, value in enumerate(values, start=1):
            ws.cell(row=first_row + i, column=col, value=value)
    last_row = first_row + len(measurements) - 1

    ws.cell(row=last_row + 2, column=1,
            value=f"Gemessen in {int(rows.sum())} von {rows.size} Bildzeilen. Ausgenommen sind Stellen, an denen "
                  "im ersten Bild keine Wunde erkannt wurde (z.B. dunkler Rand, Linie quer durch den Spalt), "
                  "plus ein Sicherheitsabstand.")
    ws.cell(row=last_row + 3, column=1,
            value="Abstand = Breite des Spalts je Bildzeile, zugewachsene Zeilen zählen als 0.")

    if measurements:
        ws.add_chart(line_chart(ws, f"{run_name}: Zuwachsen der Wunde", first_row, last_row), f"I{start}")

    return start + max(MIN_BLOCK_ROWS, len(measurements) + 6)


def save_excel(run_name: str, image_paths: list, measurements: list, rows: np.ndarray, target_dir: Path) -> Path:
    #Einzelner Versuch: eine Datei im _marked-Ordner
    wb = Workbook()
    ws = wb.active
    ws.title = "Auswertung"
    write_run(ws, 1, run_name, image_paths, measurements, rows)
    set_widths(ws, [11, 22, 14, 12, 10, 13, 16])

    out = Path(target_dir) / f"{run_name}_auswertung.xlsx"
    wb.save(out)
    return out


def write_average(ws, start: int, results: list) -> int:
    #Mittelwert und Standardabweichung je Zeitpunkt ueber alle Versuche.
    #Versuche mit weniger Zeitpunkten zaehlen nur bei den vorhandenen mit.
    ws.cell(row=start, column=1, value="Durchschnitt über alle Versuche").font = Font(bold=True, size=14)

    header = [
        "Zeitpunkt", "Anzahl Versuche",
        "Fläche relativ (%) Mittelwert", "Fläche relativ (%) Std.-Abw.",
        "Abstand Mittelwert (px) Mittelwert", "Abstand Mittelwert (px) Std.-Abw.",
        "Abstand Max (px) Mittelwert", "Abstand Max (px) Std.-Abw.",
        # Hilfsspalten fuer das Band im Graphen
        "Band unten (Mittelwert − Std.-Abw.)", "Band Breite (2 × Std.-Abw.)",
    ]
    write_header(ws, start + 1, header)

    n_points = max((len(m) for _, _, m, _ in results), default=0)

    def mean_sd(values: list) -> tuple:
        values = [v for v in values if v is not None]
        if not values:
            return None, None
        sd = float(np.std(values, ddof=1)) if len(values) > 1 else 0.0
        return round(float(np.mean(values)), 2), round(sd, 2)

    first_row = start + 2
    for t in range(n_points):
        present = [m for _, _, m, _ in results if len(m) > t]
        rel = [relative_areas(m)[t] for _, _, m, _ in results if len(m) > t]

        rel_mean, rel_sd = mean_sd(rel)
        row = [t + 1, len(present), rel_mean, rel_sd]
        for key in ("mean", "max"):
            row.extend(mean_sd([m[t][key] for m in present]))

        if rel_mean is None:
            row.extend([None, None])
        else:
            row.extend([round(rel_mean - rel_sd, 2), round(2 * rel_sd, 2)])

        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=first_row + t, column=col, value=value)
            if col >= BAND_COLUMN:
                cell.font = Font(color="808080")

    if n_points:
        ws.add_chart(average_chart(ws, first_row, first_row + n_points - 1), f"L{start}")

    return start + max(MIN_BLOCK_ROWS + 3, n_points + 3)


def average_chart(ws, first_row: int, last_row: int):
    #Mittelwert als Linie, Standardabweichung als Band dahinter.
    #Das Band ist eine gestapelte Flaeche: unsichtbarer Sockel bis
    #Mittelwert - SD, darauf 2*SD in hellblau.
    categories = Reference(ws, min_col=1, min_row=first_row, max_row=last_row)

    band = AreaChart()
    band.grouping = "stacked"
    band.add_data(Reference(ws, min_col=BAND_COLUMN, max_col=BAND_COLUMN + 1, min_row=first_row - 1, max_row=last_row), titles_from_data=True)
    band.set_categories(categories)

    sockel, breite = band.series
    sockel.graphicalProperties.noFill = True
    sockel.graphicalProperties.line.noFill = True
    breite.graphicalProperties.solidFill = BAND_COLOR
    breite.graphicalProperties.line.noFill = True
    breite.tx = SeriesLabel(v="± Standardabweichung")

    line = LineChart()
    line.add_data(Reference(ws, min_col=3, min_row=first_row - 1, max_row=last_row), titles_from_data=True)
    line.set_categories(categories)
    mittel = line.series[0]
    mittel.graphicalProperties.line.solidFill = LINE_COLOR
    mittel.graphicalProperties.line.width = 28575
    mittel.marker = Marker(symbol="circle", size=7)
    mittel.marker.graphicalProperties.solidFill = LINE_COLOR
    mittel.marker.graphicalProperties.line.solidFill = LINE_COLOR
    mittel.smooth = False
    mittel.tx = SeriesLabel(v="Mittelwert")

    band.title = "Durchschnitt: Zuwachsen der Wunde (± Standardabweichung)"
    style_axes(band, last_row - first_row + 1)
    band += line
    band.legend.legendEntry = [LegendEntry(idx=0, delete=True)]  # Sockel nicht in der Legende
    band.height = 11
    return band


def save_batch_excel(results: list, out_path: Path) -> Path:
    #results: Liste von (Versuchsname, Bildpfade, Messungen, Messbereich)
    #Alle Versuche untereinander auf einem Blatt, ganz unten der Durchschnitt.
    wb = Workbook()
    ws = wb.active
    ws.title = "Auswertung"

    row = 1
    for run_name, image_paths, measurements, rows in results:
        row = write_run(ws, row, run_name, image_paths, measurements, rows)

    write_average(ws, row + 1, results)
    set_widths(ws, [11, 22, 14, 14, 14, 14, 13, 13, 16, 14])

    wb.save(out_path)
    return Path(out_path)
